# -*- coding: utf-8 -*-
"""
HTML v7 -> редактируемый PPTX.
Текст и фигуры — родные объекты PowerPoint; иконки, графика SVG-диаграмм и
кот вставляются картинками. Геометрия снимается настоящим Chromium.
"""
from pathlib import Path
import sys
from playwright.sync_api import sync_playwright

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE

BASE = Path(__file__).resolve().parent
URL = "http://127.0.0.1:8777/setevye_informatsionnye_resursy_visual_mockup_v7.html"
IMG = BASE / "_imgs"
OUT = BASE / "presentation_editable.pptx"
SCALE = 2  # множитель растра для чётких иконок

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "start": PP_ALIGN.LEFT,
         "justify": PP_ALIGN.JUSTIFY, "": PP_ALIGN.LEFT}


def capture():
    IMG.mkdir(exist_ok=True)
    js = (BASE / "extract.js").read_text(encoding="utf-8")
    with sync_playwright() as p:
        br = p.chromium.launch(headless=True)
        page = br.new_page(viewport={"width": 1400, "height": 1000},
                           device_scale_factor=SCALE)
        page.goto(URL, wait_until="networkidle")
        page.wait_for_timeout(600)
        data = page.evaluate(js)

        # погасить текст диаграмм и подпись кота, чтобы не дублировался с нативным
        page.add_style_tag(content="svg:not(.lucide) text{visibility:hidden!important}"
                                   ".cat-caption{visibility:hidden!important}")
        page.wait_for_timeout(100)

        for s in data["slides"]:
            for im in s["images"]:
                loc = page.locator(f'[data-cap="{im["id"]}"]')
                loc.screenshot(path=str(IMG / f'{im["id"]}.png'))
        br.close()
    return data


def add_box(slide, b, epx):
    x, y, w, h = (Emu(round(b[k] * epx)) for k in ("x", "y", "w", "h"))
    rad = b.get("radius", 0)
    if rad == -1:
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    elif rad and rad > 0.5:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        try:
            shp.adjustments[0] = min(0.5, rad / min(b["w"], b["h"]))
        except Exception:
            pass
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    if b.get("fill"):
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(b["fill"])
    else:
        shp.fill.background()
    if b.get("stroke"):
        shp.line.color.rgb = RGBColor.from_string(b["stroke"])
        shp.line.width = Emu(max(1, round(b.get("strokeW", 1) * epx)))
    else:
        shp.line.fill.background()
    return shp


def add_text(slide, t, epx):
    # перенос только для реально многострочного текста; короткие метки — одной строкой
    multiline = ("\n" not in t["text"]) and (t["h"] > t["size"] * 1.6)
    pad_x, pad_y = (8, 2)
    align = t.get("align", "left")
    left = t["x"] - pad_x / 2
    width = t["w"] + pad_x
    if align in ("right", "end"):
        left = t["x"] - pad_x
    top = t["y"] - pad_y / 2
    height = t["h"] + pad_y

    tb = slide.shapes.add_textbox(Emu(round(left * epx)), Emu(round(top * epx)),
                                  Emu(round(width * epx)), Emu(round(height * epx)))
    tf = tb.text_frame
    tf.word_wrap = multiline
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    size = t["size"]
    ratio = (t.get("lh", size * 1.2) / size) if size else 1.0
    lines = t["text"].split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
        if ratio > 1.05:
            p.line_spacing = ratio
        run = p.add_run()
        run.text = line
        f = run.font
        f.name = "Arial"
        f.size = Pt(size * 0.75)
        f.bold = bool(t.get("bold"))
        f.color.rgb = RGBColor.from_string(t.get("color", "000000"))
        sp = t.get("spacing", 0)
        if sp and sp > 0.1:
            run._r.get_or_add_rPr().set("spc", str(int(sp * 0.75 * 100)))


def build(data):
    epx = data["epx"]
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]

    for s in data["slides"]:
        slide = prs.slides.add_slide(blank)
        for b in s["boxes"]:
            add_box(slide, b, epx)
        for im in s["images"]:
            path = IMG / f'{im["id"]}.png'
            if path.exists():
                slide.shapes.add_picture(str(path), Emu(round(im["x"] * epx)),
                                         Emu(round(im["y"] * epx)),
                                         Emu(round(im["w"] * epx)),
                                         Emu(round(im["h"] * epx)))
        for t in s["texts"]:
            add_text(slide, t, epx)

    prs.save(str(OUT))
    print("Готово:", OUT)
    print("Слайдов:", len(data["slides"]))


if __name__ == "__main__":
    data = capture()
    build(data)
